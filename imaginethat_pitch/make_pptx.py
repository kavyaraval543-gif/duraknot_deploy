import re
from pptx import Presentation
from pptx.util import Emu
md=open("SCRIPT.md").read()
notes=[re.sub(r"^> ?","",b.strip(),flags=re.M).replace("**","") for b in re.findall(r"### Slide \d.*?\n(.*?)(?=\n### |\n---)",md,re.S)]
p=Presentation(); p.slide_width=Emu(12192000); p.slide_height=Emu(6858000)
for i in range(8):
    s=p.slides.add_slide(p.slide_layouts[6])
    s.shapes.add_picture(f"slides/slide_{i+1}.png",0,0,p.slide_width,p.slide_height)
    s.notes_slide.notes_text_frame.text=notes[i]
p.save("Pulse_ImagineThat_Slides.pptx"); print(len(notes))
